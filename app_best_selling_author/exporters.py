"""
exporters.py
Universal Multi-Format Export Engine
Exports stories and screenplays to:
- Microsoft Word (.docx)
- Publication-Grade PDF (.pdf) with Screenplay Courier or Prose Serif styling
- Standard Markdown (.md)
- Industry Fountain (.fountain) for Final Draft / Highland
"""

import io
import re
from typing import Dict, Any

def get_text_metrics(text: str) -> Dict[str, Any]:
    """Calculates word count, character count, reading time, and estimated screenplay pages."""
    words = len(text.split())
    chars = len(text)
    # Average adult silent reading speed: ~225 words per minute
    reading_time_mins = max(1, round(words / 225))
    # Screenplay rule of thumb: ~200-250 words per page (1 page = 1 minute screen time)
    estimated_pages = max(1, round(words / 230))
    
    return {
        "words": words,
        "chars": chars,
        "reading_time_mins": reading_time_mins,
        "estimated_screen_mins": estimated_pages
    }

def export_to_markdown(text: str) -> bytes:
    """Exports raw content as UTF-8 encoded markdown bytes."""
    return text.encode("utf-8")

def export_to_fountain(text: str) -> bytes:
    """
    Exports screenplay content to .fountain industry format.
    Strips unnecessary markdown embellishments while retaining scene sluglines and cues.
    """
    lines = text.splitlines()
    fountain_lines = []
    for line in lines:
        stripped = line.strip()
        # Remove markdown headers like ## or ###
        cleaned = re.sub(r'^#+\s*', '', stripped)
        fountain_lines.append(cleaned)
    
    content = "\n".join(fountain_lines)
    return content.encode("utf-8")

def export_to_docx(text: str, title: str = "Bestseller Manuscript", is_screenplay: bool = False) -> bytes:
    """Generates a professionally styled Microsoft Word document (.docx)."""
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    
    doc = Document()
    
    # Page Margins
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1.2 if is_screenplay else 1)
        section.right_margin = Inches(1)
    
    # Configure base font
    style = doc.styles['Normal']
    font = style.font
    if is_screenplay:
        font.name = 'Courier New'
        font.size = Pt(12)
        font.color.rgb = RGBColor(20, 20, 20)
    else:
        font.name = 'Georgia'
        font.size = Pt(11.5)
        font.color.rgb = RGBColor(30, 30, 30)

    lines = text.split("\n")
    for line in lines:
        line_str = line.strip()
        if not line_str:
            doc.add_paragraph("")
            continue
            
        if line_str.startswith("# "):
            h = doc.add_heading(line_str.lstrip("# ").strip(), level=1)
            h.alignment = WD_ALIGN_PARAGRAPH.CENTER if not is_screenplay else WD_ALIGN_PARAGRAPH.LEFT
        elif line_str.startswith("## "):
            doc.add_heading(line_str.lstrip("# ").strip(), level=2)
        elif line_str.startswith("### "):
            doc.add_heading(line_str.lstrip("# ").strip(), level=3)
        elif line_str.startswith("---") or line_str.startswith("==="):
            p = doc.add_paragraph("________________________________________")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        else:
            p = doc.add_paragraph()
            # Check for screenplay character cue or slugline
            if is_screenplay:
                if re.match(r'^(?:INT\.|EXT\.|INT\./EXT\.|EXT\./INT\.)', line_str):
                    run = p.add_run(line_str.upper())
                    run.bold = True
                elif line_str.isupper() and len(line_str) < 35 and not line_str.endswith("."):
                    # Character cue
                    p.paragraph_format.left_indent = Inches(2.0)
                    run = p.add_run(line_str)
                    run.bold = True
                elif line_str.startswith("(") and line_str.endswith(")"):
                    # Parenthetical
                    p.paragraph_format.left_indent = Inches(1.5)
                    run = p.add_run(line_str)
                    run.italic = True
                else:
                    # Action or regular dialogue
                    p.add_run(line_str)
            else:
                p.add_run(line_str)
                p.paragraph_format.line_spacing = 1.25

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()

def export_to_pdf(text: str, title: str = "Bestseller Manuscript", is_screenplay: bool = False) -> bytes:
    """Generates a publication-grade PDF using ReportLab with custom typography."""
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
    from xml.sax.saxutils import escape

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        leftMargin=72,
        rightMargin=72,
        topMargin=72,
        bottomMargin=72
    )

    styles = getSampleStyleSheet()
    
    if is_screenplay:
        font_name = "Courier"
        font_bold = "Courier-Bold"
        body_size = 10.5
        leading = 13.5
    else:
        font_name = "Times-Roman"
        font_bold = "Times-Bold"
        body_size = 11.0
        leading = 15.0

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName=font_bold,
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=12,
        alignment=1 if not is_screenplay else 0
    )

    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading2'],
        fontName=font_bold,
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#0F172A'),
        spaceBefore=14,
        spaceAfter=6
    )

    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName=font_name,
        fontSize=body_size,
        leading=leading,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=8
    )

    slug_style = ParagraphStyle(
        'DocSlug',
        parent=body_style,
        fontName=font_bold,
        fontSize=body_size,
        leading=leading,
        spaceBefore=10,
        spaceAfter=4
    )

    character_style = ParagraphStyle(
        'DocChar',
        parent=body_style,
        fontName=font_bold,
        fontSize=body_size,
        leading=leading,
        leftIndent=140,
        spaceBefore=8,
        spaceAfter=2
    )

    dialogue_style = ParagraphStyle(
        'DocDial',
        parent=body_style,
        fontName=font_name,
        fontSize=body_size,
        leading=leading,
        leftIndent=80,
        rightIndent=80,
        spaceAfter=6
    )

    story = []
    lines = text.split("\n")
    
    for line in lines:
        line_str = line.strip()
        if not line_str:
            story.append(Spacer(1, 6))
            continue

        safe_text = escape(line_str)

        if line_str.startswith("# "):
            story.append(Paragraph(escape(line_str.lstrip("# ").strip()), title_style))
        elif line_str.startswith("## "):
            story.append(Paragraph(escape(line_str.lstrip("# ").strip()), h2_style))
        elif line_str.startswith("### "):
            story.append(Paragraph(f"<b>{escape(line_str.lstrip('# ').strip())}</b>", h2_style))
        elif line_str.startswith("---") or line_str.startswith("==="):
            story.append(HRFlowable(width="100%", thickness=0.8, color=colors.HexColor('#CBD5E1'), spaceBefore=8, spaceAfter=8))
        else:
            if is_screenplay:
                if re.match(r'^(?:INT\.|EXT\.|INT\./EXT\.|EXT\./INT\.)', line_str):
                    story.append(Paragraph(safe_text.upper(), slug_style))
                elif line_str.isupper() and len(line_str) < 35 and not line_str.endswith("."):
                    story.append(Paragraph(safe_text, character_style))
                elif line_str.startswith("(") and line_str.endswith(")"):
                    story.append(Paragraph(f"<i>{safe_text}</i>", dialogue_style))
                else:
                    story.append(Paragraph(safe_text, body_style))
            else:
                story.append(Paragraph(safe_text, body_style))

    doc.build(story)
    buffer.seek(0)
    return buffer.getvalue()
